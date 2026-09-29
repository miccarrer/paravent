"""Regenerate the test documents. Every name, number and address is invented.

    uv run python tests/fixtures/make_fixtures.py [/path/to/LiberationSerif-Regular.ttf]

Without the font, the scanned letter and the photo are left as they are.
The outputs are committed, so the tests need neither this script nor the font.
"""

import random
import sys
from datetime import datetime, timezone
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path

import docx

from PIL import Image, ImageDraw, ImageFilter, ImageFont

HERE = Path(__file__).parent

LETTER = [
    "Caisse d'Allocations Familiales de Exempleville",
    "Madame Élodie Lefèvre-Garçon",
    "12, rue des Écoles — 99000 Exempleville",
    "Objet : réexamen de vos droits à l'aide au logement",
    "Madame, nous accusons réception de votre dossier n° 1234567 A.",
    "Vos prestations d'août sont maintenues à 312,45 € par mois.",
    "Tél. : 01 23 45 67 89 — courriel : elodie.lefevre@exemple.fr",
    "Numéro de sécurité sociale : 2 84 05 99 123 456 78",
    "IBAN : FR76 3000 6000 0112 3456 7890 189",
    "Veuillez agréer, Madame, l'expression de nos salutations distinguées.",
]

NATIVE = [
    "Pôle Exemple - Agence de Exempleville",
    "Convocation à un entretien le 12 février à 9 h 30.",
    "Merci de vous présenter muni de votre pièce d'identité.",
]

ATTACHMENT = [
    "Pôle Exemple - Agence de Exempleville",
    "Convocation : atelier collectif le 12 février à 14 h.",
    "Cet atelier est obligatoire.",
]


def scanned_letter(font_path: str) -> Image.Image:
    width, height = 1654, 1100  # A4 width at 200 dpi, top of the page only
    image = Image.new("L", (width, height), 255)
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(font_path, 38)
    for index, line in enumerate(LETTER):
        draw.text((140, 120 + 70 * index), line, font=font, fill=20)
    # Look like a scan: slight skew, blur and speckles.
    rng = random.Random(1)
    image = image.rotate(0.8, fillcolor=255, resample=Image.Resampling.BICUBIC)
    image = image.filter(ImageFilter.GaussianBlur(0.8))
    pixels = image.load()
    for _ in range(8000):
        pixels[rng.randrange(width), rng.randrange(height)] = rng.choice((0, 90, 200))
    return image


def native_pdf(lines: list[str]) -> bytes:
    """A minimal one-page PDF with a real text layer (Helvetica, WinAnsi)."""

    def pdf_string(text: str) -> str:
        raw = text.encode("cp1252")
        return "".join(chr(b) if 32 <= b < 127 and chr(b) not in "()\\" else f"\\{b:03o}" for b in raw)

    ops = ["BT", "/F1 14 Tf", "72 770 Td", "18 TL"]
    ops += [f"({pdf_string(line)}) Tj T*" for line in lines]
    ops.append("ET")
    stream = "\n".join(ops).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(out)


def word_letter() -> docx.Document:
    document = docx.Document()
    document.core_properties.created = document.core_properties.modified = datetime(2026, 1, 1)
    document.core_properties.author = "Exemple"
    document.add_heading("Bail d'habitation", level=1)
    document.add_paragraph("Entre les soussignés :")
    document.add_paragraph("Monsieur Gérard Exemplaire, bailleur", style="List Bullet")
    document.add_paragraph("Madame Élodie Lefèvre-Garçon, locataire", style="List Bullet")
    document.add_heading("Loyer", level=2)
    paragraph = document.add_paragraph("Le loyer mensuel est fixé à ")
    paragraph.add_run("540,00 €").bold = True
    paragraph.add_run(", charges comprises.")
    table = document.add_table(rows=3, cols=2)
    for row, (label, value) in zip(table.rows, [("Poste", "Montant"), ("Loyer", "480,00 €"), ("Charges", "60,00 €")]):
        row.cells[0].text, row.cells[1].text = label, value
    document.sections[0].header.paragraphs[0].text = "Agence Exemple Immobilier — 99000 Exempleville"
    return document


def email_message() -> EmailMessage:
    message = EmailMessage()
    message["From"] = "Conseillère Pôle Exemple <conseil@pole-exemple.fr>"
    message["To"] = "Élodie Lefèvre-Garçon <elodie.lefevre@exemple.fr>"
    message["Subject"] = "Votre convocation du 12 février"
    message["Date"] = format_datetime(datetime(2026, 2, 2, 10, 15, tzinfo=timezone.utc))
    message.set_content("Bonjour Madame,\n\nVous trouverez ci-joint votre convocation.\n\nCordialement,\nVotre conseillère")
    message.add_alternative("<p>Bonjour Madame,</p><p>Vous trouverez <b>ci-joint</b> votre convocation.</p>"
                            "<p>Cordialement,<br>Votre conseillère</p>", subtype="html")
    message.add_attachment(native_pdf(ATTACHMENT), maintype="application",
                           subtype="pdf", filename="convocation.pdf")
    return message


def main() -> None:
    if len(sys.argv) > 1:
        scan = scanned_letter(sys.argv[1])
        scan.convert("RGB").save(HERE / "courrier-scanne.pdf", resolution=200, quality=70)
        scan.crop((100, 80, 1300, 380)).save(HERE / "courrier-photo.png", optimize=True)
    (HERE / "convocation-native.pdf").write_bytes(native_pdf(NATIVE))
    word_letter().save(HERE / "bail.docx")
    (HERE / "convocation.eml").write_bytes(bytes(email_message()))


if __name__ == "__main__":
    main()
